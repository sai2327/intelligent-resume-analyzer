"""DOCX text extraction using python-docx. Accepts a file path or an in-memory file-like object."""

from docx import Document
from docx.opc.exceptions import PackageNotFoundError


class DOCXExtractionError(Exception):
    """Raised when a DOCX file cannot be read or contains no extractable text."""


def extract_text_from_docx(file) -> str:
    """Extract all paragraph and table text from a DOCX. `file` may be a path
    string or a file-like object (e.g. an uploaded file's stream)."""
    try:
        document = Document(file)
    except PackageNotFoundError as exc:
        raise DOCXExtractionError(
            f"The DOCX file is corrupt or not a valid Word document: {exc}"
        ) from exc
    except Exception as exc:
        raise DOCXExtractionError(f"Could not open DOCX file: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Headers and footers: resumes frequently park contact info or the name in
    # the document header, which plain paragraph iteration misses entirely.
    seen_blocks = set()
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                text = paragraph.text.strip()
                if text and text not in seen_blocks:
                    seen_blocks.add(text)
                    parts.append(text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    full_text = "\n".join(parts).strip()
    if not full_text:
        raise DOCXExtractionError("No extractable text was found in this DOCX file.")
    return full_text
