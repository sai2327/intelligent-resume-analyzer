"""Post-process the built deck: within each <a:p>, keep only the first <a:pPr>.
pptxgenjs emits one <a:pPr> per run; OOXML allows at most one per paragraph, and
PowerPoint re-lays-out affected paragraphs (mixed bold/regular lead-ins)."""
import sys
from pptx import Presentation
from pptx.oxml.ns import qn

path = sys.argv[1]
prs = Presentation(path)
fixed = 0
for slide in prs.slides:
    for p in slide.shapes._spTree.iter(qn("a:p")):
        pprs = p.findall(qn("a:pPr"))
        for extra in pprs[1:]:
            p.remove(extra)
            fixed += 1
prs.save(path)
print(f"removed {fixed} duplicate a:pPr elements")
